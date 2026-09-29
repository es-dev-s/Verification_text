"""
Extract text from .docx: paragraphs, tables, and OCR of embedded images when text is sparse.
"""

from __future__ import annotations

import os
import tempfile
from typing import Any

from extractor.image_extract import ocr_image_file

MIN_TEXT_CHARS = 40


def extract_docx(path: str, *, filename: str | None = None) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is not installed. Run: pip install python-docx"
        ) from exc

    doc = Document(path)
    parts: list[str] = []
    warnings: list[str] = []

    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        parts.append(f"--- table {table_index} ---")
        for row in table.rows:
            cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    body = "\n".join(parts).strip()
    method = "docx"
    ocr_confidence: float | None = None

    if len(body) < MIN_TEXT_CHARS:
        image_texts: list[str] = []
        confs: list[float] = []
        try:
            image_parts = _iter_docx_images(doc)
            for index, (ext, blob) in enumerate(image_parts, start=1):
                with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                    tmp.write(blob)
                    tmp_path = tmp.name
                try:
                    spans, avg = ocr_image_file(tmp_path, page_number=index)
                    text = "\n".join(s["text"] for s in spans if s.get("text")).strip()
                    if text:
                        image_texts.append(f"--- image {index} ---\n{text}")
                    if avg is not None:
                        confs.append(avg)
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"image {index} OCR failed: {exc}")
                finally:
                    try:
                        os.remove(tmp_path)
                    except OSError:
                        pass
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"docx image OCR skipped: {exc}")

        if image_texts:
            body = (body + "\n\n" if body else "") + "\n\n".join(image_texts)
            method = "docx+ocr" if body[: len(body)] else "ocr"
            ocr_confidence = round(sum(confs) / len(confs), 1) if confs else None

    return {
        "ok": bool(body.strip()),
        "filename": filename,
        "raw_text": body,
        "method": method,
        "ocr_confidence": ocr_confidence,
        "page_count": 1,
        "pages": [
            {
                "page": 1,
                "method": method,
                "source": method,
                "ocr_confidence": ocr_confidence,
                "char_count": len(body),
            }
        ],
        "warnings": warnings,
    }


def _iter_docx_images(doc: Any) -> list[tuple[str, bytes]]:
    """Pull embedded image blobs from the docx package."""
    images: list[tuple[str, bytes]] = []
    try:
        package = doc.part.package
        for rel in package.parts:
            content_type = getattr(rel, "content_type", "") or ""
            if not str(content_type).startswith("image/"):
                continue
            blob = rel.blob
            if not blob:
                continue
            ext = ".png"
            if "jpeg" in content_type or "jpg" in content_type:
                ext = ".jpg"
            elif "gif" in content_type:
                ext = ".gif"
            elif "bmp" in content_type:
                ext = ".bmp"
            images.append((ext, blob))
    except Exception:
        # Fallback: relationships on the main document part
        try:
            for rel in doc.part.rels.values():
                if "image" not in str(getattr(rel, "reltype", "")):
                    continue
                target = rel.target_part
                blob = target.blob
                ctype = getattr(target, "content_type", "") or ""
                ext = ".jpg" if "jpeg" in ctype or "jpg" in ctype else ".png"
                images.append((ext, blob))
        except Exception:
            return images
    return images
